import json
import os
from pathlib import Path

import openpyxl
from loguru import logger


DEFAULT_DSN = "postgresql://spider_xhs:spider_xhs_password@127.0.0.1:15432/spider_xhs"


NOTE_COLUMNS = [
    "note_id", "note_url", "note_type", "user_id", "home_url", "nickname",
    "avatar", "title", "desc", "liked_count", "collected_count",
    "comment_count", "share_count", "video_cover", "video_addr",
    "image_list", "tags", "upload_time", "ip_location", "keyword", "raw_json",
    "comments_crawl_status", "comments_collected_count", "comments_target_count",
    "comments_next_cursor", "comments_has_more",
]

COMMENT_COLUMNS = [
    "comment_id", "note_id", "parent_comment_id", "note_title", "content",
    "user_id", "user_nickname", "like_count", "create_time", "ip_location",
    "is_hot", "sub_comment_count", "raw_json",
]


SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS xhs_crawl_tasks (
    task_id BIGSERIAL PRIMARY KEY,
    keyword TEXT NOT NULL,
    status TEXT NOT NULL,
    config JSONB NOT NULL DEFAULT '{}'::jsonb,
    note_count INTEGER NOT NULL DEFAULT 0,
    comment_count INTEGER NOT NULL DEFAULT 0,
    error_message TEXT,
    started_at TIMESTAMPTZ NOT NULL DEFAULT now(),
    finished_at TIMESTAMPTZ
);

CREATE TABLE IF NOT EXISTS xhs_notes (
    note_id TEXT PRIMARY KEY,
    note_url TEXT,
    note_type TEXT,
    user_id TEXT,
    home_url TEXT,
    nickname TEXT,
    avatar TEXT,
    title TEXT,
    "desc" TEXT,
    liked_count INTEGER NOT NULL DEFAULT 0,
    collected_count INTEGER NOT NULL DEFAULT 0,
    comment_count INTEGER NOT NULL DEFAULT 0,
    share_count INTEGER NOT NULL DEFAULT 0,
    video_cover TEXT,
    video_addr TEXT,
    image_list JSONB NOT NULL DEFAULT '[]'::jsonb,
    tags JSONB NOT NULL DEFAULT '[]'::jsonb,
    upload_time TEXT,
    ip_location TEXT,
    keyword TEXT,
    raw_json JSONB NOT NULL DEFAULT '{}'::jsonb,
    comments_crawl_status TEXT NOT NULL DEFAULT 'pending',
    comments_collected_count INTEGER NOT NULL DEFAULT 0,
    comments_target_count INTEGER NOT NULL DEFAULT 0,
    comments_next_cursor TEXT NOT NULL DEFAULT '',
    comments_has_more BOOLEAN NOT NULL DEFAULT true,
    comments_updated_at TIMESTAMPTZ NOT NULL DEFAULT now(),
    crawled_at TIMESTAMPTZ NOT NULL DEFAULT now(),
    updated_at TIMESTAMPTZ NOT NULL DEFAULT now()
);

ALTER TABLE xhs_notes ADD COLUMN IF NOT EXISTS comments_crawl_status TEXT NOT NULL DEFAULT 'pending';
ALTER TABLE xhs_notes ADD COLUMN IF NOT EXISTS comments_collected_count INTEGER NOT NULL DEFAULT 0;
ALTER TABLE xhs_notes ADD COLUMN IF NOT EXISTS comments_target_count INTEGER NOT NULL DEFAULT 0;
ALTER TABLE xhs_notes ADD COLUMN IF NOT EXISTS comments_next_cursor TEXT NOT NULL DEFAULT '';
ALTER TABLE xhs_notes ADD COLUMN IF NOT EXISTS comments_has_more BOOLEAN NOT NULL DEFAULT true;
ALTER TABLE xhs_notes ADD COLUMN IF NOT EXISTS comments_updated_at TIMESTAMPTZ NOT NULL DEFAULT now();

CREATE TABLE IF NOT EXISTS xhs_comments (
    comment_id TEXT PRIMARY KEY,
    note_id TEXT NOT NULL REFERENCES xhs_notes(note_id) ON DELETE CASCADE,
    parent_comment_id TEXT,
    note_title TEXT,
    content TEXT,
    user_id TEXT,
    user_nickname TEXT,
    like_count INTEGER NOT NULL DEFAULT 0,
    create_time TEXT,
    ip_location TEXT,
    is_hot BOOLEAN NOT NULL DEFAULT false,
    sub_comment_count INTEGER NOT NULL DEFAULT 0,
    raw_json JSONB NOT NULL DEFAULT '{}'::jsonb,
    crawled_at TIMESTAMPTZ NOT NULL DEFAULT now(),
    updated_at TIMESTAMPTZ NOT NULL DEFAULT now()
);

CREATE INDEX IF NOT EXISTS idx_xhs_notes_keyword ON xhs_notes(keyword);
CREATE INDEX IF NOT EXISTS idx_xhs_notes_user_id ON xhs_notes(user_id);
CREATE INDEX IF NOT EXISTS idx_xhs_comments_note_id ON xhs_comments(note_id);
CREATE INDEX IF NOT EXISTS idx_xhs_comments_parent ON xhs_comments(parent_comment_id);
"""


def _to_int(value):
    if value in (None, ""):
        return 0
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def normalize_note_row(note, keyword="", raw=None):
    return {
        "note_id": note.get("note_id", ""),
        "note_url": note.get("note_url", ""),
        "note_type": note.get("note_type", ""),
        "user_id": note.get("user_id", ""),
        "home_url": note.get("home_url", ""),
        "nickname": note.get("nickname", ""),
        "avatar": note.get("avatar", ""),
        "title": note.get("title", ""),
        "desc": note.get("desc", ""),
        "liked_count": _to_int(note.get("liked_count")),
        "collected_count": _to_int(note.get("collected_count")),
        "comment_count": _to_int(note.get("comment_count")),
        "share_count": _to_int(note.get("share_count")),
        "video_cover": note.get("video_cover"),
        "video_addr": note.get("video_addr"),
        "image_list": note.get("image_list") or [],
        "tags": note.get("tags") or [],
        "upload_time": note.get("upload_time", ""),
        "ip_location": note.get("ip_location", ""),
        "keyword": keyword,
        "raw_json": raw or {},
        "comments_crawl_status": note.get("comments_crawl_status", "pending"),
        "comments_collected_count": _to_int(note.get("comments_collected_count")),
        "comments_target_count": _to_int(note.get("comments_target_count")),
        "comments_next_cursor": note.get("comments_next_cursor", ""),
        "comments_has_more": bool(note.get("comments_has_more", True)),
    }


def normalize_comment_row(comment, note_id, note_title="", parent_comment_id=None, raw=None):
    return {
        "comment_id": comment.get("comment_id") or comment.get("id", ""),
        "note_id": note_id,
        "parent_comment_id": parent_comment_id,
        "note_title": note_title,
        "content": comment.get("content", ""),
        "user_id": comment.get("user_id", ""),
        "user_nickname": comment.get("user_nickname", ""),
        "like_count": _to_int(comment.get("like_count")),
        "create_time": comment.get("create_time", ""),
        "ip_location": comment.get("ip_location", ""),
        "is_hot": bool(comment.get("is_hot", False)),
        "sub_comment_count": _to_int(comment.get("sub_comment_count")),
        "raw_json": raw or {},
    }


def _excel_cell(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _jsonb(value):
    from psycopg.types.json import Jsonb
    return Jsonb(value)


def _adapt_json_fields(row, fields):
    adapted = dict(row)
    for field in fields:
        adapted[field] = _jsonb(adapted.get(field))
    return adapted


def export_rows_to_excel(rows, file_path, columns):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(columns)
    for row in rows:
        ws.append([_excel_cell(row.get(column, "")) for column in columns])
    Path(file_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(file_path)
    logger.info(f"Excel 已导出: {file_path}")


class PostgresStorage:
    def __init__(self, dsn=None):
        self.dsn = self.resolve_dsn(dsn, os.environ)
        try:
            import psycopg
        except ImportError as exc:
            raise ImportError("缺少 psycopg，请先运行: pip install -r requirements.txt") from exc
        self._psycopg = psycopg
        self.conn = psycopg.connect(self.dsn)
        self.conn.autocommit = True

    @staticmethod
    def resolve_dsn(dsn=None, environ=None):
        environ = environ or os.environ
        return dsn or environ.get("DATABASE_URL") or DEFAULT_DSN

    def close(self):
        self.conn.close()

    def init_schema(self):
        with self.conn.cursor() as cur:
            cur.execute(SCHEMA_SQL)

    def start_task(self, keyword, config):
        with self.conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO xhs_crawl_tasks (keyword, status, config)
                VALUES (%s, %s, %s)
                RETURNING task_id
                """,
                (keyword, "running", _jsonb(config)),
            )
            return cur.fetchone()[0]

    def finish_task(self, task_id, status, note_count=0, comment_count=0, error_message=None):
        with self.conn.cursor() as cur:
            cur.execute(
                """
                UPDATE xhs_crawl_tasks
                SET status=%s, note_count=%s, comment_count=%s, error_message=%s, finished_at=now()
                WHERE task_id=%s
                """,
                (status, note_count, comment_count, error_message, task_id),
            )

    def upsert_note(self, note, keyword="", raw=None):
        row = _adapt_json_fields(normalize_note_row(note, keyword, raw), ["image_list", "tags", "raw_json"])
        with self.conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO xhs_notes (
                    note_id, note_url, note_type, user_id, home_url, nickname,
                    avatar, title, "desc", liked_count, collected_count,
                    comment_count, share_count, video_cover, video_addr,
                    image_list, tags, upload_time, ip_location, keyword, raw_json
                )
                VALUES (
                    %(note_id)s, %(note_url)s, %(note_type)s, %(user_id)s, %(home_url)s, %(nickname)s,
                    %(avatar)s, %(title)s, %(desc)s, %(liked_count)s, %(collected_count)s,
                    %(comment_count)s, %(share_count)s, %(video_cover)s, %(video_addr)s,
                    %(image_list)s, %(tags)s, %(upload_time)s, %(ip_location)s, %(keyword)s, %(raw_json)s
                )
                ON CONFLICT (note_id) DO UPDATE SET
                    note_url=EXCLUDED.note_url,
                    note_type=EXCLUDED.note_type,
                    user_id=EXCLUDED.user_id,
                    home_url=EXCLUDED.home_url,
                    nickname=EXCLUDED.nickname,
                    avatar=EXCLUDED.avatar,
                    title=EXCLUDED.title,
                    "desc"=EXCLUDED."desc",
                    liked_count=EXCLUDED.liked_count,
                    collected_count=EXCLUDED.collected_count,
                    comment_count=EXCLUDED.comment_count,
                    share_count=EXCLUDED.share_count,
                    video_cover=EXCLUDED.video_cover,
                    video_addr=EXCLUDED.video_addr,
                    image_list=EXCLUDED.image_list,
                    tags=EXCLUDED.tags,
                    upload_time=EXCLUDED.upload_time,
                    ip_location=EXCLUDED.ip_location,
                    keyword=EXCLUDED.keyword,
                    raw_json=EXCLUDED.raw_json,
                    updated_at=now()
                """,
                row,
            )

    def upsert_comment(self, comment, note_id, note_title="", parent_comment_id=None, raw=None):
        row = _adapt_json_fields(
            normalize_comment_row(comment, note_id, note_title, parent_comment_id, raw),
            ["raw_json"],
        )
        if not row["comment_id"]:
            return
        with self.conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO xhs_comments (
                    comment_id, note_id, parent_comment_id, note_title, content,
                    user_id, user_nickname, like_count, create_time, ip_location,
                    is_hot, sub_comment_count, raw_json
                )
                VALUES (
                    %(comment_id)s, %(note_id)s, %(parent_comment_id)s, %(note_title)s, %(content)s,
                    %(user_id)s, %(user_nickname)s, %(like_count)s, %(create_time)s, %(ip_location)s,
                    %(is_hot)s, %(sub_comment_count)s, %(raw_json)s
                )
                ON CONFLICT (comment_id) DO UPDATE SET
                    note_id=EXCLUDED.note_id,
                    parent_comment_id=EXCLUDED.parent_comment_id,
                    note_title=EXCLUDED.note_title,
                    content=EXCLUDED.content,
                    user_id=EXCLUDED.user_id,
                    user_nickname=EXCLUDED.user_nickname,
                    like_count=EXCLUDED.like_count,
                    create_time=EXCLUDED.create_time,
                    ip_location=EXCLUDED.ip_location,
                    is_hot=EXCLUDED.is_hot,
                    sub_comment_count=EXCLUDED.sub_comment_count,
                    raw_json=EXCLUDED.raw_json,
                    updated_at=now()
                """,
                row,
            )

    def count_comments(self, note_id):
        with self.conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) FROM xhs_comments WHERE note_id=%s", (note_id,))
            return cur.fetchone()[0]

    def update_comment_progress(self, note_id, status, collected_count, target_count, next_cursor, has_more):
        with self.conn.cursor() as cur:
            cur.execute(
                """
                UPDATE xhs_notes
                SET comments_crawl_status=%s,
                    comments_collected_count=%s,
                    comments_target_count=%s,
                    comments_next_cursor=%s,
                    comments_has_more=%s,
                    comments_updated_at=now(),
                    updated_at=now()
                WHERE note_id=%s
                """,
                (status, collected_count, target_count, next_cursor or "", bool(has_more), note_id),
            )

    def get_comment_progress_by_note_ids(self, note_ids):
        if not note_ids:
            return {}
        placeholders = ", ".join(["%s"] * len(note_ids))
        with self.conn.cursor() as cur:
            cur.execute(
                f"""
                SELECT note_id, comments_crawl_status, comments_collected_count,
                       comments_target_count, comments_next_cursor, comments_has_more
                FROM xhs_notes
                WHERE note_id IN ({placeholders})
                """,
                list(note_ids),
            )
            return {
                row[0]: {
                    "comments_crawl_status": row[1],
                    "comments_collected_count": row[2],
                    "comments_target_count": row[3],
                    "comments_next_cursor": row[4],
                    "comments_has_more": row[5],
                }
                for row in cur.fetchall()
            }

    def fetch_notes(self, keyword=None):
        sql = "SELECT " + ", ".join([f'"{c}"' if c == "desc" else c for c in NOTE_COLUMNS]) + " FROM xhs_notes"
        params = []
        if keyword:
            sql += " WHERE keyword=%s"
            params.append(keyword)
        sql += " ORDER BY crawled_at DESC"
        with self.conn.cursor() as cur:
            cur.execute(sql, params)
            return [dict(zip(NOTE_COLUMNS, row)) for row in cur.fetchall()]

    def fetch_comments(self, keyword=None):
        sql = """
            SELECT c.comment_id, c.note_id, c.parent_comment_id, c.note_title, c.content,
                   c.user_id, c.user_nickname, c.like_count, c.create_time, c.ip_location,
                   c.is_hot, c.sub_comment_count, c.raw_json
            FROM xhs_comments c
        """
        params = []
        if keyword:
            sql += " JOIN xhs_notes n ON n.note_id = c.note_id WHERE n.keyword=%s"
            params.append(keyword)
        sql += " ORDER BY c.crawled_at DESC"
        with self.conn.cursor() as cur:
            cur.execute(sql, params)
            return [dict(zip(COMMENT_COLUMNS, row)) for row in cur.fetchall()]

    def export_to_excel(self, output_dir, keyword=None):
        Path(output_dir).mkdir(parents=True, exist_ok=True)
        prefix = keyword or "xhs"
        note_file = Path(output_dir) / f"{prefix}_笔记数据_数据库导出.xlsx"
        comment_file = Path(output_dir) / f"{prefix}_评论数据_数据库导出.xlsx"
        export_rows_to_excel(self.fetch_notes(keyword), note_file, NOTE_COLUMNS)
        export_rows_to_excel(self.fetch_comments(keyword), comment_file, COMMENT_COLUMNS)
        return note_file, comment_file
